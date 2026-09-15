import os
import sys
import json
import csv
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Setup Django environment
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'hint_system.settings')

import django
django.setup()

from hints.models import Attempt, Problem, UserProfile, HintDelivery

def build_attempt_text(att, step_num, total_steps_in_prefix):
    lines = []
    lines.append(f"--- STEP {step_num} of {total_steps_in_prefix} ---")
    lines.append(f"Timestamp: {att.created_at.strftime('%Y-%m-%d %H:%M:%S UTC')}")
    lines.append(f"Action / Source: {att.source}")
    lines.append(f"Status: {att.status}")
    if att.execution_time:
        lines.append(f"Execution Time: {att.execution_time}")
    
    if att.evaluation_details:
        eval_str = json.dumps(att.evaluation_details, indent=2)
        lines.append(f"Evaluation Details: {eval_str}")
    
    lines.append("Code:")
    lines.append(att.code if att.code else "[NO CODE SUBMITTED]")
    lines.append("")
    return "\n".join(lines)

def generate_dataset():
    print("=" * 80)
    print("GENERATING HUMAN ANNOTATION DATASET EXCEL WORKBOOK")
    print("=" * 80)

    # Fetch attempts grouped by valid session_id
    attempts = list(Attempt.objects.select_related('problem').all().order_by('created_at'))
    
    sessions_map = {}
    for att in attempts:
        sid = att.session_id.strip() if att.session_id else ""
        # Filtering Rule 1 & 2: Must have valid session_id (exclude orphaned legacy sessions)
        if not sid or sid.startswith("LEGACY_ORPHANED"):
            continue
        if sid not in sessions_map:
            sessions_map[sid] = []
        sessions_map[sid].append(att)

    # Filtering Rule 3: Exclude isolated single-attempt sessions (num_attempts < 2)
    # Rule 4: DO NOT exclude multi-attempt sessions merely because code did not change.
    valid_sessions_map = {}
    for sid, att_list in sessions_map.items():
        if len(att_list) >= 2:
            valid_sessions_map[sid] = att_list

    print(f"Total Valid Multi-Attempt Sessions Identified: {len(valid_sessions_map)}")

    wb = openpyxl.Workbook()
    # Default sheet
    ws_data = wb.active
    ws_data.title = "Human Annotation Dataset"

    # Define Styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Navy
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    empty_label_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Soft Amber for annotation cells
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "Session ID",
        "Student",
        "Problem",
        "Prefix Name",
        "Attempt History (Trajectory up to Prefix)",
        "Human Label",
        "Evidence",
        "Confidence",
        "Annotator"
    ]

    ws_data.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws_data.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    ws_data.row_dimensions[1].height = 28

    total_attempts_included = 0
    total_prefixes_created = 0
    students_represented = set()
    problems_represented = set()
    session_summary_rows = []

    row_num = 2
    for sid, att_list in valid_sessions_map.items():
        att_list.sort(key=lambda x: x.created_at)
        n_atts = len(att_list)
        total_attempts_included += n_atts

        student_id = att_list[0].user_id or "UNKNOWN"
        students_represented.add(student_id)

        problem_title = att_list[0].problem.title if att_list[0].problem else "UNKNOWN"
        problem_id = att_list[0].problem.problem_id if att_list[0].problem else "UNKNOWN"
        problems_represented.add(problem_id)

        session_prefixes_count = n_atts - 1

        session_summary_rows.append({
            'session_id': sid,
            'student': student_id,
            'problem': f"{problem_title} ({problem_id})",
            'total_attempts': n_atts,
            'prefixes_generated': session_prefixes_count,
            'duration_secs': (att_list[-1].created_at - att_list[0].created_at).total_seconds()
        })

        # Generate prefixes A1..A2, A1..A3, ..., A1..An
        for prefix_idx in range(2, n_atts + 1):
            prefix_attempts = att_list[:prefix_idx]
            prefix_name = f"Prefix {prefix_idx - 1} (Steps 1..{prefix_idx})"
            
            history_blocks = []
            for idx, a in enumerate(prefix_attempts, 1):
                history_blocks.append(build_attempt_text(a, idx, prefix_idx))
            
            full_history_text = "\n".join(history_blocks)

            row_data = [
                sid,
                student_id,
                f"{problem_title} [{problem_id}]",
                prefix_name,
                full_history_text,
                "", # Human Label (EMPTY FOR ANNOTATORS)
                "", # Evidence (EMPTY FOR ANNOTATORS)
                "", # Confidence (EMPTY FOR ANNOTATORS)
                ""  # Annotator (EMPTY FOR ANNOTATORS)
            ]

            ws_data.append(row_data)

            # Apply formatting to data row
            ws_data.cell(row=row_num, column=1).alignment = Alignment(vertical="top")
            ws_data.cell(row=row_num, column=2).alignment = Alignment(vertical="top")
            ws_data.cell(row=row_num, column=3).alignment = Alignment(vertical="top")
            ws_data.cell(row=row_num, column=4).alignment = Alignment(vertical="top", horizontal="center")
            ws_data.cell(row=row_num, column=5).alignment = Alignment(vertical="top", wrap_text=True)

            # Highlight empty annotation columns with soft amber fill
            for c in range(6, 10):
                cell = ws_data.cell(row=row_num, column=c)
                cell.fill = empty_label_fill
                cell.alignment = Alignment(vertical="top")
                cell.border = border_thin

            for c in range(1, 10):
                ws_data.cell(row=row_num, column=c).border = border_thin

            ws_data.row_dimensions[row_num].height = min(max(80, len(full_history_text.splitlines()) * 12), 400)
            row_num += 1
            total_prefixes_created += 1

    # Adjust Column Widths for Sheet 1
    col_widths = {
        1: 36, # Session ID
        2: 32, # Student
        3: 32, # Problem
        4: 22, # Prefix Name
        5: 75, # Attempt History
        6: 18, # Human Label
        7: 30, # Evidence
        8: 14, # Confidence
        9: 16  # Annotator
    }
    for col_idx, width in col_widths.items():
        ws_data.column_dimensions[get_column_letter(col_idx)].width = width

    # =========================================================================
    # Sheet 2: Annotation Guidelines & README
    # =========================================================================
    ws_readme = wb.create_sheet(title="Annotation Guidelines & README")
    ws_readme.views.sheetView[0].showGridLines = True

    readme_content = [
        ("HUMAN ANNOTATION PROTOCOL & GUIDELINES", "H1"),
        ("Research Dataset: Hint-Timing Adaptive Scaffolding Project", "H2"),
        ("", "NORMAL"),
        ("1. OVERVIEW & OBJECTIVE", "H2"),
        ("The purpose of this dataset is to conduct independent human annotation on student problem-solving trajectories to evaluate student struggle dynamics. Annotators will evaluate trajectory prefixes and categorize student behavior into PRODUCTIVE, UNPRODUCTIVE, or UNCERTAIN struggle.", "NORMAL"),
        ("", "NORMAL"),
        ("2. KEY DEFINITIONS", "H2"),
        ("• PRODUCTIVE STRUGGLE:", "BOLD"),
        ("  The student is actively making meaningful progress toward understanding or solving the problem. Evidence includes evolving error types, improving test outputs, refining logic, exploring logical alternatives, or effectively applying hint guidance.", "NORMAL"),
        ("• UNPRODUCTIVE STRUGGLE:", "BOLD"),
        ("  The student is stuck in a repetitive, unhelpful, or stagnant loop without structural progress. Evidence includes repeated identical submissions, unchanged broken logic across multiple attempts, thrashing/guessing code without hypothesis, or ignoring critical errors.", "NORMAL"),
        ("• UNCERTAIN:", "BOLD"),
        ("  The available trajectory prefix lacks sufficient information or evidence to confidently determine whether the struggle is productive or unproductive.", "NORMAL"),
        ("", "NORMAL"),
        ("3. CORE ANNOTATION PRINCIPLE: JUDGE PROCESS, NOT OUTCOME", "H2"),
        ("⚠️ CRITICAL INSTRUCTION: Annotators must evaluate the student's PROBLEM-SOLVING PROCESS up to the current prefix step, NOT whether the student eventually solved or failed the problem.", "WARNING"),
        ("• Do NOT assume 'Eventually Solved = Productive'. A student may stumble into a pass by random guessing or copy-pasting.", "NORMAL"),
        ("• Do NOT assume 'Failed = Unproductive'. A student may be engaged in highly productive conceptual exploration and hypothesis testing even if their time ran out before final completion.", "NORMAL"),
        ("", "NORMAL"),
        ("4. HOW TO JUDGE A TRAJECTORY PREFIX", "H2"),
        ("• Each row presents a TRAJECTORY PREFIX (e.g. Steps 1..3).", "NORMAL"),
        ("• You must judge the student's behavior ONLY using the steps provided in that prefix.", "NORMAL"),
        ("• STRICT RULE: Do NOT look ahead to future attempts beyond the prefix. In real-time tutoring systems, future attempts are unavailable at decision time.", "NORMAL"),
        ("", "NORMAL"),
        ("5. REPEATED & ZERO-CODE-CHANGE ATTEMPTS", "H2"),
        ("• Multi-attempt sessions with repeated code or zero edit diffs are INTENTIONALLY INCLUDED in this dataset.", "NORMAL"),
        ("• Repeated identical submissions provide strong empirical evidence of unproductive struggle (stagnation/hesitation) and must be evaluated accordingly.", "NORMAL"),
        ("", "NORMAL"),
        ("6. INDEPENDENT ANNOTATION REQUIREMENTS", "H2"),
        ("• All three annotators must work completely INDEPENDENTLY without consulting each other or viewing each other's labels.", "NORMAL"),
        ("• Do not discuss specific session rows until all individual annotations are finalized and submitted for inter-annotator agreement scoring.", "NORMAL"),
        ("", "NORMAL"),
        ("7. ALLOWED VALUES & FIELD SPECIFICATIONS", "H2"),
        ("• Human Label (Must be one of):", "BOLD"),
        ("  - PRODUCTIVE", "NORMAL"),
        ("  - UNPRODUCTIVE", "NORMAL"),
        ("  - UNCERTAIN", "NORMAL"),
        ("", "NORMAL"),
        ("• Evidence Guidance (Select or specify the primary indicator):", "BOLD"),
        ("  - Meaningful code change", "NORMAL"),
        ("  - Error evolution", "NORMAL"),
        ("  - Test/result improvement", "NORMAL"),
        ("  - Change of approach", "NORMAL"),
        ("  - Repeated error", "NORMAL"),
        ("  - Repeated approach/code", "NORMAL"),
        ("  - No meaningful progress", "NORMAL"),
        ("  - Exploration without clear progress", "NORMAL"),
        ("  - Insufficient evidence", "NORMAL"),
        ("", "NORMAL"),
        ("• Confidence (Integer rating):", "BOLD"),
        ("  - 1 = Low Confidence", "NORMAL"),
        ("  - 2 = Medium Confidence", "NORMAL"),
        ("  - 3 = High Confidence", "NORMAL"),
        ("", "NORMAL"),
        ("• Annotator (Your assigned name / identifier):", "BOLD"),
        ("  - e.g. Annotator_A, Annotator_B, Annotator_C", "NORMAL")
    ]

    h1_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    h1_font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    h2_font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
    bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    warning_font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")
    norm_font = Font(name="Segoe UI", size=10, color="1E293B")

    for idx, (text, style_type) in enumerate(readme_content, 1):
        cell = ws_readme.cell(row=idx, column=1, value=text)
        if style_type == "H1":
            cell.fill = h1_fill
            cell.font = h1_font
            ws_readme.row_dimensions[idx].height = 30
        elif style_type == "H2":
            cell.font = h2_font
        elif style_type == "BOLD":
            cell.font = bold_font
        elif style_type == "WARNING":
            cell.font = warning_font
        else:
            cell.font = norm_font

    ws_readme.column_dimensions['A'].width = 110

    # =========================================================================
    # Sheet 3: Valid Sessions Inventory
    # =========================================================================
    ws_inv = wb.create_sheet(title="Valid Sessions Inventory")
    inv_headers = ["Session ID", "Student User", "Problem Title & ID", "Total Attempts", "Prefixes Created", "Session Duration (s)"]
    ws_inv.append(inv_headers)

    for col_idx in range(1, len(inv_headers) + 1):
        cell = ws_inv.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_inv.row_dimensions[1].height = 25

    for r_idx, summary in enumerate(session_summary_rows, 2):
        ws_inv.append([
            summary['session_id'],
            summary['student'],
            summary['problem'],
            summary['total_attempts'],
            summary['prefixes_generated'],
            summary['duration_secs']
        ])
        for c in range(1, 7):
            ws_inv.cell(row=r_idx, column=c).border = border_thin

    inv_widths = {1: 36, 2: 32, 3: 45, 4: 16, 5: 18, 6: 20}
    for col_idx, width in inv_widths.items():
        ws_inv.column_dimensions[get_column_letter(col_idx)].width = width

    # Save Excel Workbook
    output_excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'human_annotation_dataset.xlsx')
    wb.save(output_excel_path)
    print(f"Successfully created Human Annotation Dataset Excel workbook at: {output_excel_path}")

    # Summary Report Metrics
    print("\n" + "=" * 80)
    print("HUMAN ANNOTATION DATASET GENERATION SUMMARY")
    print("=" * 80)
    print(f"Number of Valid Sessions Included: {len(valid_sessions_map)}")
    print(f"Number of Unique Students Represented: {len(students_represented)}")
    print(f"Number of Unique Problems Represented: {len(problems_represented)}")
    print(f"Total Attempts Included: {total_attempts_included}")
    print(f"Total Trajectory Prefixes Created: {total_prefixes_created}")
    print(f"Output File Location: {output_excel_path}")
    print("=" * 80)

    return {
        'valid_sessions_count': len(valid_sessions_map),
        'students_count': len(students_represented),
        'problems_count': len(problems_represented),
        'total_attempts_included': total_attempts_included,
        'total_prefixes_created': total_prefixes_created,
        'output_excel_path': output_excel_path
    }

if __name__ == '__main__':
    generate_dataset()
